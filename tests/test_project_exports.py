"""导出使用纯内存样本、模拟查询和 SQL 编译，不访问业务数据库。"""

import json
from datetime import date, datetime
from decimal import Decimal
from importlib import import_module
from io import BytesIO
from types import SimpleNamespace as NS
from uuid import uuid4

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql
from sqlalchemy.orm import Session

import workflow_models  # noqa: F401 注册 ORM 关联
from routers import annotation_projects, interpretation_projects, recruitment_projects  # noqa: F401 注册全部关联模型
import project_export_routes as routes
import project_export_service as service
from database import get_db


def project(**extra):
    return NS(**{
        "id": uuid4(), "order_no": "0000123", "project_name": "示例项目",
        "project_status": "in_progress", "client_full_name": "客户公司",
        "client_name": "客户公司",
        "client_short_name": "客户", "client_code": "00001", "client_id": uuid4(),
        "customer_consultation_time": datetime(2026, 10, 1, 23, 59, 59),
        "parent_project_id": None, "role_assignments": [],
        "time_ranges": [], "language_directions": [], "interpreter_assignments": [],
        "language_items": [], "price_items": [], "assignees": [],
        "candidates": [], "progress_records": [], **extra,
    })


def workbook(module, mode, items, **kwargs):
    return load_workbook(BytesIO(service.projects_to_xlsx(module, mode, [items], **kwargs)))


def row(sheet, number=2):
    return dict(zip((cell.value for cell in sheet[1]), (cell.value for cell in sheet[number])))


def parser(module):
    router = import_module(f"routers.{module}_projects")
    return lambda raw, db: router._field_filters(raw, db) if module == "annotation" else router._field_filters(raw)


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_typed_project_exports_and_safe_text(module):
    p = project(project_name="=HYPERLINK(\"bad\")\x01", service_fee_amount=Decimal("1234.56"))
    book = workbook(module, "projects", [p, p])
    sheet = book["母订单" if module == "annotation" else "项目"]
    assert sheet.max_row == 2  # 按 ID 去重
    values = row(sheet)
    assert values["订单号"] == "0000123"
    assert values["项目名称"].startswith("'=HYPERLINK")
    assert "\x01" not in values["项目名称"]
    assert isinstance(values["客户咨询时间"], datetime)
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref.endswith("2")
    if module == "recruitment":
        assert values["固定服务费"] == 1234.56


def test_interpretation_relations_are_separate_and_keep_multilingual_direction():
    p = project(time_ranges=[NS(sequence_no=1, scheduled_start=datetime(2026, 10, 1), scheduled_end=datetime(2026, 10, 2))],
                language_directions=[NS(sequence_no=1, display="中文 ↔ 英语 ↔ 日语", required_count=3)],
                interpreter_assignments=[NS(translator_name="张三", translator_code="001", customer_rating="satisfied")])
    book = workbook("interpretation", "projects", [p])
    assert book.sheetnames == ["项目", "时间安排", "语言方向", "译员安排"]
    assert row(book["语言方向"])["语言方向"] == "中文 ↔ 英语 ↔ 日语"
    assert row(book["译员安排"])["客户对译员评价"] == "满意"
    assert row(book["时间安排"])["预定结束时间"] == datetime(2026, 10, 2)


def test_annotation_parent_child_rates_and_custom_fields():
    p = project(project_status="project_in_progress", custom_values={"field-id": "原值"})
    child = project(parent_project_id=p.id, parent_order_no=p.order_no, order_no="AP-001-C01",
                    price_items=[NS(amount=Decimal("2.125"), currency="USD", unit="条")],
                    assignees=[NS(person_name="人员一", assignment_role="quality_inspector", assignment_status="completed",
                                  language_item=NS(display="英语"), audio_duration_value=Decimal("12.5"), audio_duration_unit="minute",
                                  rate=NS(amount=Decimal("0.025"), currency="USD", unit="second"))])
    book = workbook("annotation", "projects", [p, child, child], custom_labels={"field-id": "业务标签"})
    assert book.sheetnames == ["母订单", "子订单", "语言项", "客户单价", "人员安排"]
    assert book["母订单"].max_row == book["子订单"].max_row == 2
    assert row(book["子订单"])["母订单号"] == p.order_no
    assert json.loads(row(book["母订单"])["自定义字段"]) == {"业务标签": "原值"}
    assert row(book["客户单价"])["客户单价"] == 2.125
    values = row(book["人员安排"])
    assert (values["人员单价"], values["工作量单位"], values["人员计价单位"], values["安排角色"]) == (0.025, "分钟", "秒", "质检员")
    assert "密码" not in values


def test_recruitment_export_contains_every_followup_without_cartesian_duplicates():
    candidate = NS(id=uuid4(), candidate_name="李四", stage="onboarded", actual_onboard_date=date(2026, 10, 8),
                   communications=[NS(sequence_no=1, communication_date=date(2026, 10, 7), details="已联系"),
                                   NS(sequence_no=2, communication_date=date(2026, 10, 8), details="确认入职")],
                   interviews=[NS(round_no=1, interview_date=date(2026, 10, 6), details="面试通过")])
    p = project(candidates=[candidate], progress_records=[NS(note="项目跟进", to_status="probation")],
                language_directions=[NS(direction_type="single", label="英语")])
    book = workbook("recruitment", "projects", [p])
    assert book.sheetnames == ["项目", "语言需求", "项目进度", "候选人", "候选人沟通", "面试记录"]
    assert book["候选人"].max_row == 2
    assert book["候选人沟通"].max_row == 3
    assert book["面试记录"].max_row == 2
    assert row(book["候选人"])["候选人阶段"] == "已入职"
    assert row(book["项目进度"])["变更后项目状态"] == "已入职保用期"
    tracking = workbook("recruitment", "candidates", [p])
    assert tracking.sheetnames == ["候选人", "候选人沟通", "面试记录"]
    assert row(tracking["候选人"])["实际入职日期"] == datetime(2026, 10, 8)
    assert all("应付金额" not in [cell.value for cell in sheet[1]] for sheet in tracking)


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_real_model_customer_properties_are_exported(module):
    from models import Client, SubClient
    model = getattr(import_module(f"{module}_models"), f"{module.title()}Project")
    p = model(id=uuid4(), order_no="001", project_status="in_progress")
    p.client = Client(client_name="母公司", client_short_name="母简称", client_code="P001")
    p.sub_client = SubClient(client_name="子公司", client_short_name="子简称", sub_client_code="S001")
    values = row(workbook(module, "projects", [p])["母订单" if module == "annotation" else "项目"])
    assert values["客户全称"] == "子公司"
    assert values["客户编号"] == "S001"
    assert values["客户简称"] == "子简称"


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_customer_reconciliation_keeps_raw_values_in_pending_without_invented_settlement(module):
    p = project(customer_budget="预算人民币2万左右", service_fee_type="monthly_salary_multiple", service_fee_multiplier=Decimal("1.5"),
                price_items=[NS(amount=Decimal("3"), currency="CNY", unit="小时")])
    book = workbook(module, "reconciliation", [p])
    assert book.sheetnames == ["对账单", "待补数据"]
    assert book["对账单"].max_row == 1
    values = row(book["待补数据"])
    assert "账单月份" in values["待补原因"]
    assert "已确认结算金额" in values["待补原因"]
    assert "结算总额" not in values
    if module == "interpretation":
        assert values["客户预算原文"] == "预算人民币2万左右"
    elif module == "annotation":
        assert values["客户单价"] == 3
    else:
        assert values["服务费倍数"] == 1.5
        assert values["服务费类型"] == "月薪倍数"


def test_annotation_reconciliation_keeps_orders_without_quotes_and_each_quote():
    book = workbook("annotation", "reconciliation", [project(), project(price_items=[NS(amount=1, currency="CNY"), NS(amount=2, currency="USD")])])
    assert book["待补数据"].max_row == 4
    assert "客户单价" in row(book["待补数据"])["待补原因"]


@pytest.mark.parametrize("module,relation", [("interpretation", "interpreter_assignments"), ("annotation", "assignees")])
def test_personnel_excludes_cancelled_projects_and_assignments(module, relation):
    p = project(**{relation: [NS(translator_name="当前译员", person_name="当前人员", assignment_status="assigned"),
                              NS(translator_name="取消译员", person_name="取消人员", assignment_status="cancelled")]})
    cancelled = project(project_status="cancelled", **{relation: [NS(assignment_status="assigned")]})
    book = workbook(module, "personnel", [p, cancelled])
    assert book["待补数据"].max_row == 2
    assert "当前" in (row(book["待补数据"]).get("人员姓名") or row(book["待补数据"]).get("译员姓名"))
    with pytest.raises(service.ProjectExportEmptyError):
        workbook(module, "personnel", [cancelled])


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_empty_and_each_sheet_limit(module):
    with pytest.raises(service.ProjectExportEmptyError):
        workbook(module, "projects", [])
    with pytest.raises(service.ProjectExportLimitError):
        workbook(module, "projects", [project(), project()], max_rows_per_sheet=1)
    with pytest.raises(service.ProjectExportLimitError):
        workbook(module, "reconciliation", [project(), project()], max_rows_per_sheet=1)


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_filter_override_keeps_other_filters_and_sort(module):
    original = {"created_at": {"op": "between", "from": "2020-01-01", "to": "2020-02-01"},
                "order_no": {"op": "contains", "value": "001"}}
    result = routes.build_export_filters(module, parser(module), object(), time_field="created_at",
                                         date_start=date(2026, 10, 1), date_end=date(2026, 10, 9),
                                         field_filters=json.dumps(original), keyword="项目", sort="latest_progress_desc" if module == "annotation" else None)
    assert result["keyword"] == "项目"
    assert result["field_filters"]["created_at"]["from"] == "2026-10-01"
    assert result["field_filters"]["order_no"] == original["order_no"]
    if module == "annotation":
        assert result["sort"] == "latest_progress_desc"


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_customer_id_mode_ignores_time_and_list_filters(module):
    client = uuid4()
    filters = routes.build_export_filters(module, parser(module), object(), client_id=client,
                                          time_field="bad", field_filters="bad", keyword="other", sort="bad", order_scope="child", parent_project_id=uuid4())
    assert filters == ({"client_id": client, "order_scope": "all"} if module == "annotation" else {"client_id": client})


@pytest.mark.parametrize("kwargs", [
    {"time_field": "bad"}, {"date_start": None}, {"date_end": date(2026, 9, 1)},
    {"field_filters": '{"password":{"op":"contains","value":"x"}}'},
    {"sort": "bad"},
])
def test_invalid_export_conditions_are_rejected(kwargs):
    args = {"time_field": "created_at", "date_start": date(2026, 10, 1), "date_end": date(2026, 10, 9), **kwargs}
    with pytest.raises(HTTPException) as error:
        routes.build_export_filters("interpretation", parser("interpretation"), object(), **args)
    assert error.value.status_code == 422


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_date_boundaries_and_exact_parent_client_use_existing_query(module):
    model = getattr(import_module(f"{module}_models"), f"{module.title()}Project")
    apply = import_module(f"{module}_service")._apply_filters
    client_id = uuid4()
    with Session() as db:
        query = apply(db.query(model), client_id=client_id, field_filters={"created_at": {"op": "between", "from": "2026-10-01", "to": "2026-10-09"}})
        compiled = query.statement.compile(dialect=postgresql.dialect())
    params = list(compiled.params.values())
    assert client_id in params
    assert datetime(2026, 10, 1) in params
    assert datetime(2026, 10, 10) in params  # 结束日包含整天，次日零点不包含。
    assert f"{module}_project.client_id =" in str(compiled)


def test_interpretation_scheduled_date_uses_range_overlap():
    from interpretation_service import _apply_filters
    from interpretation_models import InterpretationProject
    with Session() as db:
        query = _apply_filters(db.query(InterpretationProject), field_filters={"scheduled_date": {"op": "between", "from": "2026-10-01", "to": "2026-10-09"}})
        compiled = query.statement.compile(dialect=postgresql.dialect())
    assert "scheduled_end >=" in str(compiled)
    assert "scheduled_start <=" in str(compiled)
    assert datetime(2026, 10, 9, 23, 59, 59, 999999) in compiled.params.values()


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_batch_loading_exports_beyond_first_page_and_loads_relations(monkeypatch, module):
    calls = []
    data = [project(candidates=[NS(candidate_name=f"候选人{i}", communications=[], interviews=[])]) for i in range(503)]
    def loader(_db, **kwargs):
        calls.append(kwargs)
        return data[kwargs["skip"]:kwargs["skip"] + kwargs["limit"]]
    monkeypatch.setattr(import_module(f"{module}_service"), f"get_{module}_projects", loader)
    filters = {"keyword": "当前筛选", **({"order_scope": "child"} if module == "annotation" else {})}
    book = load_workbook(BytesIO(service.projects_to_xlsx(module, "projects", service.iter_project_batches(object(), module, filters))))
    assert sum(book[title].max_row - 1 for title in (["母订单", "子订单"] if module == "annotation" else ["项目"])) == 503
    assert [call["skip"] for call in calls] == [0, 500]
    assert all(call["keyword"] == "当前筛选" for call in calls)
    if module != "interpretation":
        assert all(call["extra_options"] for call in calls)


def test_annotation_parent_batch_includes_all_children_and_child_scope_does_not(monkeypatch):
    import annotation_service
    parent = project()
    child = project(parent_project_id=parent.id)
    monkeypatch.setattr(annotation_service, "get_annotation_projects", lambda *_a, **_k: [parent])
    class Query:
        def options(self, *_a): return self
        def filter(self, *_a): return self
        def order_by(self, *_a): return self
        def offset(self, skip): self.skip = skip; return self
        def limit(self, *_a): return self
        def all(self): return [child] if self.skip == 0 else []
    class DB:
        def query(self, *_a): return Query()
    assert list(service.iter_project_batches(DB(), "annotation", {"order_scope": "parent"})) == [[parent], [child]]
    assert list(service.iter_project_batches(object(), "annotation", {"order_scope": "child"})) == [[parent]]


@pytest.mark.parametrize("module", service.MODULE_LABELS)
def test_real_router_protects_static_export_paths_and_returns_download(monkeypatch, module):
    router = import_module(f"routers.{module}_projects").router
    paths = [route.path for route in router.routes]
    export_path = f"/projects/{module}/export"
    assert paths.index(export_path) < paths.index(f"/projects/{module}/{{project_id}}")
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_db] = lambda: object()
    client = TestClient(app)
    assert client.get(export_path, params={"time_field": "created_at", "date_start": "2026-10-01", "date_end": "2026-10-09"}).status_code in {401, 403}
    app.dependency_overrides[router.dependencies[0].dependency] = lambda: None
    seen = []
    def create(_db, mod, mode, filters):
        seen.append((mod, mode, filters))
        return service.projects_to_xlsx(mod, mode, [[project()]])
    monkeypatch.setattr(routes, "create_project_export", create)
    response = client.get(export_path, params={"time_field": "created_at", "date_start": "2026-10-01", "date_end": "2026-10-09"})
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert load_workbook(BytesIO(response.content)).sheetnames
    assert seen[0][:2] == (module, "projects")
    assert client.get(export_path, params={"time_field": "created_at", "date_start": "bad", "date_end": "2026-10-09"}).status_code == 422
    client_id = uuid4()
    response = client.get(f"/projects/{module}/reconciliation-export", params={"client_id": str(client_id)})
    assert response.status_code == 200
    assert seen[-1][2]["client_id"] == client_id


@pytest.mark.parametrize("error,status", [(service.ProjectExportEmptyError("没有数据"), 404), (service.ProjectExportLimitError("超限"), 422)])
def test_route_business_errors(monkeypatch, error, status):
    def fail(*_a, **_k): raise error
    monkeypatch.setattr(routes, "create_project_export", fail)
    with pytest.raises(HTTPException) as result:
        routes.export_response(object(), "interpretation", "projects", {}, time_field="created_at", date_start=date(2026, 10, 1), date_end=date(2026, 10, 9))
    assert result.value.status_code == status
