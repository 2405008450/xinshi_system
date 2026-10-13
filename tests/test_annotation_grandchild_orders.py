"""三级订单、片区筛选和继承保护，全部运行在临时 PostgreSQL 中。"""
from datetime import datetime
from pathlib import Path
from uuid import uuid4
from types import SimpleNamespace

import pytest
from sqlalchemy import text
from annotation_models import AnnotationArea
from annotation_ops_models import AnnotationCustomFieldDefinition, AnnotationProjectStatusHistory
from annotation_schemas import AnnotationProjectCreate
import annotation_service as service
from test_annotation_child_orders import sessions, context, item, update_payload, add_test_language


def chain(context):
    db, root_id, language_id = context
    area = AnnotationArea(label='片区'+uuid4().hex[:8], name_key=uuid4().hex)
    db.add(area); db.commit()
    child = service.create_annotation_children(db, root_id, [item(language_id)], None, uuid4().hex)[0]
    grandchild = service.create_annotation_children(db, child.id, [item(language_id, area_id=area.id)], None, uuid4().hex)[0]
    return db, root_id, child.id, grandchild.id, area.id


def save(db, project_id, **changes):
    project = service.get_annotation_project(db, project_id)
    return service.update_annotation_project(db, project_id, update_payload(project, **changes))


def test_levels_numbering_and_leaf_limit(context):
    db, root_id, child_id, leaf_id, area_id = chain(context)
    root, child, leaf = [service.get_annotation_project(db, id) for id in (root_id, child_id, leaf_id)]
    assert [root.order_level, child.order_level, leaf.order_level] == [0, 1, 2]
    assert leaf.order_no == child.order_no+'.001'
    assert root.child_count == child.child_count == 1 and leaf.child_count == 0
    with pytest.raises(ValueError, match='最多支持三层'):
        service.create_annotation_children(db, leaf.id, [item(leaf.language_items[0].source_language_id, area_id=area_id)], None, uuid4().hex)
    db.rollback()
    with pytest.raises(service.AnnotationProjectDeleteConflict):
        service.delete_annotation_project(db, child_id)
    db.rollback()
    old_no = leaf.order_no
    language_id = leaf.language_items[0].source_language_id
    service.delete_annotation_project(db, leaf_id)
    replacement = service.create_annotation_children(db, child_id, [item(language_id, area_id=area_id)], None, uuid4().hex)[0]
    assert replacement.order_no != old_no


def test_grandchild_concurrent_numbering_and_batch_rollback(context,sessions):
    from concurrent.futures import ThreadPoolExecutor
    db,root_id,child_id,leaf_id,area_id=chain(context)
    language_id=service.get_annotation_project(db,child_id).language_items[0].source_language_id
    db.rollback()
    def create(index):
        with sessions() as concurrent:
            return service.create_annotation_children(concurrent,child_id,[item(language_id,area_id=area_id)],None,f'grandchild-concurrent-{child_id}-{index}')[0].order_no
    with ThreadPoolExecutor(max_workers=4) as executor:
        numbers=list(executor.map(create,range(6)))
    assert len(set(numbers))==6
    assert service.count_annotation_projects(db,parent_project_id=child_id)==7
    with pytest.raises(ValueError,match='片区不存在'):
        service.create_annotation_children(db,child_id,[item(language_id,area_id=area_id),item(language_id,area_id=uuid4())],None,uuid4().hex)
    db.rollback()
    assert service.count_annotation_projects(db,parent_project_id=child_id)==7
    key=uuid4().hex;payload=[item(language_id,area_id=area_id)]
    first=service.create_annotation_children(db,child_id,payload,None,key)[0]
    assert service.create_annotation_children(db,child_id,payload,None,key)[0].id==first.id


def test_propagates_and_preserves_overrides_and_explicit_clear(context):
    db, root_id, child_id, leaf_id, _ = chain(context)
    save(db, root_id, potential_demand='100')
    assert service.get_annotation_project(db, leaf_id).potential_demand == '100'
    save(db, child_id, potential_demand='150')
    save(db, root_id, potential_demand='200', contact_name='新联系人')
    child, leaf = [service.get_annotation_project(db, id) for id in (child_id, leaf_id)]
    assert child.potential_demand == leaf.potential_demand == '150'
    assert leaf.contact_name == '新联系人'
    save(db, leaf_id, potential_demand=None)
    save(db, child_id, potential_demand='300')
    assert service.get_annotation_project(db, leaf_id).potential_demand is None
    leaf = save(db, leaf_id, restore_inheritance=['potential_demand'])
    assert leaf.potential_demand == '300' and 'potential_demand' not in leaf.inheritance_overrides
    save(db, child_id, restore_inheritance=['potential_demand'])
    assert service.get_annotation_project(db, leaf_id).potential_demand == '200'


def test_noop_full_save_does_not_break_inheritance(context):
    db, root_id, child_id, leaf_id, _ = chain(context)
    child = service.get_annotation_project(db, child_id)
    original = child.inheritance_overrides
    assert save(db, child_id).inheritance_overrides == original
    save(db, root_id, task_description='新任务')
    assert service.get_annotation_project(db, leaf_id).task_description == '新任务'


def test_runtime_session_preserves_inline_parent_changes_and_overrides(context, sessions):
    from sqlalchemy.orm import Session
    from annotation_inheritance_service import snapshot, finish_inline_change
    db, root_id, child_id, leaf_id, _ = chain(context)
    language_id = service.get_annotation_project(db, child_id).language_items[0].source_language_id
    sibling = service.create_annotation_children(db, root_id, [item(language_id)], None, uuid4().hex)[0]
    sibling_id = sibling.id
    db.rollback()
    # 与 database.SessionLocal 一致，验证多个下级加载不会覆盖尚未提交的上级字段。
    with Session(sessions.kw['bind'], autoflush=False) as runtime:
        parent = service.get_annotation_project(runtime, root_id)
        before = snapshot(parent)
        parent.potential_demand = '200'
        finish_inline_change(runtime, parent, before)
        runtime.commit()
        runtime.expire_all()
        assert [service.get_annotation_project(runtime, id).potential_demand for id in
                (root_id, child_id, leaf_id, sibling_id)] == ['200'] * 4
        child = service.get_annotation_project(runtime, child_id)
        before = snapshot(child)
        child.potential_demand = '150'
        finish_inline_change(runtime, child, before)
        runtime.commit()
        child = service.get_annotation_project(runtime, child_id)
        assert 'potential_demand' in child.inheritance_overrides
        parent = service.get_annotation_project(runtime, root_id)
        before = snapshot(parent)
        parent.potential_demand = '300'
        finish_inline_change(runtime, parent, before)
        runtime.commit()
        runtime.expire_all()
        assert [service.get_annotation_project(runtime, id).potential_demand for id in
                (root_id, child_id, leaf_id, sibling_id)] == ['300', '150', '150', '300']


def test_historical_orders_stay_independent_including_future_fields(context):
    db, root_id, child_id, leaf_id, _ = chain(context)
    child = service.get_annotation_project(db, child_id)
    child.inheritance_overrides = None; db.commit()
    save(db, root_id, contact_name='新母联系人')
    assert service.get_annotation_project(db, child_id).contact_name == '原联系人'
    assert save(db, child_id, potential_demand='历史自填').inheritance_overrides is None
    save(db, child_id, restore_inheritance=['contact_name'])
    field = AnnotationCustomFieldDefinition(table_code='project', field_key=uuid4().hex, field_label='新字段', data_type='text', sequence_no=900)
    db.add(field); db.commit()
    save(db, root_id, custom_values={str(field.id):'母值'}, contact_name='再次更新')
    child = service.get_annotation_project(db, child_id)
    assert child.contact_name == '再次更新' and not child.custom_values.get(str(field.id))
    assert service.get_annotation_project(db, leaf_id).contact_name == '再次更新'


def test_custom_fields_inherit_separately(context):
    db, root_id, child_id, leaf_id, _ = chain(context)
    fields = [AnnotationCustomFieldDefinition(table_code='project', field_key=uuid4().hex, field_label='字段'+str(i), data_type='text', sequence_no=910+i) for i in range(2)]
    db.add_all(fields); db.commit()
    keys=[str(field.id) for field in fields]
    save(db, root_id, custom_values={keys[0]:'A',keys[1]:'B'})
    save(db, child_id, custom_values={keys[0]:'独立',keys[1]:'B'})
    save(db, root_id, custom_values={keys[0]:'AA',keys[1]:'BB'})
    values=service.get_annotation_project(db, leaf_id).custom_values
    assert {key:values[key] for key in keys} == {keys[0]:'独立',keys[1]:'BB'}
    save(db, child_id, restore_inheritance=['custom:'+keys[0]])
    assert service.get_annotation_project(db, leaf_id).custom_values[keys[0]] == 'AA'


def test_area_rollup_scopes_and_counts(context):
    db, root_id, child_id, leaf_id, area_id = chain(context)
    filters={'area_ids':[area_id]}
    for scope, expected in [('parent',root_id),('child',child_id),('grandchild',leaf_id)]:
        rows=service.get_annotation_projects(db, order_scope=scope, **filters)
        assert [row.id for row in rows] == [expected]
        assert service.count_annotation_projects(db, order_scope=scope, **filters) == 1
        assert len(rows[0].area_labels) == 1
    assert service.get_annotation_project(db, root_id).area_labels == service.get_annotation_project(db, leaf_id).area_labels
    assert service.count_annotation_projects(db, area_ids=[uuid4()]) == 0


def test_exports_include_three_levels_and_area(context):
    from project_export_service import iter_project_batches, projects_to_xlsx
    from openpyxl import load_workbook
    from io import BytesIO
    db,root_id,child_id,leaf_id,area_id=chain(context)
    filters={'order_scope':'parent','area_ids':[area_id]}
    batches=list(iter_project_batches(db,'annotation',filters))
    assert [row.id for batch in batches for row in batch]==[root_id,child_id,leaf_id]
    book=load_workbook(BytesIO(projects_to_xlsx('annotation','projects',batches)))
    assert all(title in book.sheetnames for title in ['订单','子订单','三级订单'])
    assert '片区' in [cell.value for cell in book['三级订单'][1]]


def test_leaf_language_fixed_and_child_change_propagates(context):
    db, root_id, child_id, leaf_id, _ = chain(context)
    other=add_test_language(db)
    leaf=service.get_annotation_project(db,leaf_id)
    with pytest.raises(ValueError, match='必须与所属子订单一致'):
        save(db,leaf_id,language_items=[{'id':leaf.language_items[0].id,'source_language_id':other}])
    db.rollback()
    child=service.get_annotation_project(db,child_id)
    save(db,child_id,language_items=[{'id':child.language_items[0].id,'source_language_id':other}])
    assert service.get_annotation_project(db,leaf_id).language_items[0].source_language_id == other


def test_language_conflict_and_atomic_rollback(context):
    db, root_id, child_id, leaf_id, _ = chain(context)
    other=add_test_language(db)
    original=service.get_annotation_project(db,child_id).language_items[0].source_language_id
    save(db,leaf_id,price_items=[{'source_language_id':original,'amount':5,'unit':'条'}])
    child=service.get_annotation_project(db,child_id)
    with pytest.raises(ValueError, match='已独立维护'):
        save(db,child_id,project_name='不应保存',language_items=[{'id':child.language_items[0].id,'source_language_id':other}])
    db.rollback()
    assert service.get_annotation_project(db,child_id).language_items[0].source_language_id == original
    assert service.get_annotation_project(db,child_id).project_name != '不应保存'


def test_prices_inherit_as_group(context):
    db,root_id,child_id,leaf_id,_=chain(context)
    save(db,root_id,price_items=[{'amount':2,'unit':'条'}])
    assert service.get_annotation_project(db,leaf_id).price_items[0].amount == 2
    old_id=service.get_annotation_project(db,leaf_id).price_items[0].id
    save(db,root_id,price_items=[{'amount':3,'unit':'条'}])
    assert service.get_annotation_project(db,leaf_id).price_items[0].id == old_id
    save(db,child_id,price_items=[])
    save(db,root_id,price_items=[{'amount':4,'unit':'条'}])
    assert service.get_annotation_project(db,leaf_id).price_items == []


def test_inline_rates_and_people_preserve_child_identity(context):
    from resource_models import ResourcePerson, ResourceCapability
    from annotation_ops_schemas import AssigneeRateWrite
    from annotation_ops_service import save_assignee_rate, delete_assignee_rate
    db,root_id,child_id,leaf_id,_=chain(context)
    person=ResourcePerson(full_name='继承单价人员');db.add(person);db.flush()
    db.add(ResourceCapability(person_id=person.id,capability_type='annotation',status='active'));db.commit()
    save(db,root_id,assignees=[{'person_id':person.id,'rate':{'amount':2,'unit':'item'}}])
    leaf=service.get_annotation_project(db,leaf_id);original_id=leaf.assignees[0].id
    root=service.get_annotation_project(db,root_id)
    save_assignee_rate(db,root.assignees[0].id,AssigneeRateWrite(amount=3,unit='item'))
    assert service.get_annotation_project(db,leaf_id).assignees[0].rate.amount==3
    assert service.get_annotation_project(db,leaf_id).assignees[0].id==original_id
    child=service.get_annotation_project(db,child_id)
    delete_assignee_rate(db,child.assignees[0].id)
    save_assignee_rate(db,root.assignees[0].id,AssigneeRateWrite(amount=4,unit='item'))
    assert service.get_annotation_project(db,leaf_id).assignees[0].rate is None


def test_inline_managers_and_handover_inherit_without_stale_roles(context,monkeypatch):
    from models import AppUser
    monkeypatch.setattr('project_workbench_service.validate_assignment_map',lambda *_:None)
    db,root_id,child_id,leaf_id,_=chain(context)
    users=[AppUser(username=uuid4().hex,password_hash='isolated-test',full_name='继承负责人') for _ in range(2)]
    db.add_all(users);db.commit()
    service.update_annotation_project_managers(db,root_id,users[0].id,[users[0].id])
    for id in [child_id,leaf_id]:
        row=service.get_annotation_project(db,id)
        assert row.client_manager_id==users[0].id
        assert any(role.assignee_id==users[0].id for role in row.workbench_responsibilities)
    service.update_annotation_project_managers(db,child_id,users[0].id,[users[1].id])
    service.update_annotation_project_managers(db,root_id,users[1].id,[])
    row=service.get_annotation_project(db,leaf_id)
    assert row.client_manager_id==users[1].id and any(role.assignee_id==users[1].id for role in row.workbench_responsibilities)
    root=service.get_annotation_project(db,root_id);root.client_manager_id=users[0].id
    service._sync_child_shared_fields(db,root,users[1].id)
    db.commit()
    assert service.get_annotation_project(db,leaf_id).client_manager_id==users[0].id


def test_status_propagation_records_own_history_and_priority(context):
    db,root_id,child_id,leaf_id,_=chain(context)
    service.update_annotation_project_status(db,root_id,'project_in_progress',datetime(2026,10,12,12,0),'母单推进')
    assert service.get_annotation_project(db,leaf_id).project_status == 'project_in_progress'
    assert db.query(AnnotationProjectStatusHistory).filter_by(project_id=leaf_id,to_status='project_in_progress').count()==1
    service.update_annotation_project_status(db,child_id,'paused',datetime(2026,10,12,12,1),'子单暂停')
    service.update_annotation_project_status(db,root_id,'ended',datetime(2026,10,12,12,2),'母单结束')
    assert service.get_annotation_project(db,leaf_id).project_status=='paused'
    service.update_annotation_project_priority(db,root_id,'high')
    assert service.get_annotation_project(db,leaf_id).priority=='high'


def test_api_areas_and_grandchildren(context,monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from database import get_db
    from routers import annotation_projects as routes,auth
    db,root_id,child_id,leaf_id,area_id=chain(context)
    app=FastAPI();app.include_router(routes.router)
    app.dependency_overrides[get_db]=lambda:db
    app.dependency_overrides[auth.get_current_user]=lambda:SimpleNamespace(id=None)
    monkeypatch.setattr(auth,'user_has_permission',lambda *_:True)
    monkeypatch.setattr(auth,'get_user_permission_codes',lambda *_:{'projects:read','projects:write'})
    with TestClient(app) as client:
        label='新片区'+uuid4().hex
        first=client.post('/projects/annotation/areas',json={'label':' '+label+' '})
        assert first.status_code==200,first.text
        assert client.post('/projects/annotation/areas',json={'label':label}).json()['id']==first.json()['id']
        page=client.get('/projects/annotation/page',params={'order_scope':'grandchild','area_ids':str(area_id)})
        assert page.status_code==200,page.text
        assert page.json()['items'][0]['id']==str(leaf_id)
        assert client.get(f'/projects/annotation/{child_id}/children').json()['total']==1
        assert client.get(f'/projects/annotation/{leaf_id}/children').status_code==400


def test_migration_is_repeatable_and_preserves_legacy(context):
    db,root_id,language_id=context
    db.rollback()
    sql=(Path(__file__).resolve().parents[1]/'data/migrations/20261017_annotation_grandchild_inheritance.sql').read_text(encoding='utf-8')
    with db.get_bind().connect().execution_options(isolation_level='AUTOCOMMIT') as connection:
        connection.exec_driver_sql(sql);connection.exec_driver_sql(sql)
        assert connection.execute(text("SELECT inheritance_overrides FROM annotation_project WHERE order_no='AP-LEGACY-KEEP'")).scalar() is None
