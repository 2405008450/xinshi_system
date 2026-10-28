"""只读区域匹配：来源文字经明确去字后，必须被目标字段连续整体包含。"""
from collections import defaultdict
from io import BytesIO, StringIO
import csv
from pathlib import Path
from typing import Literal
import unicodedata
from zipfile import ZipFile, BadZipFile

from fastapi import HTTPException
from pydantic import BaseModel, Field, model_validator, ConfigDict

RegionField = Literal['ancestral_home', 'native_place', 'residence_address']
FIELDS = ('ancestral_home', 'native_place', 'residence_address')
MAX_ROWS = 1000
MAX_COLUMNS = 50
MAX_BYTES = 10 * 1024 * 1024


class RegionMapping(BaseModel):
    model_config = ConfigDict(extra='forbid')
    source_column: int = Field(ge=0, lt=MAX_COLUMNS)
    target_fields: list[RegionField] = Field(min_length=1, max_length=3)


class RegionMatchWrite(BaseModel):
    model_config = ConfigDict(extra='forbid')
    rows: list[list[str]] = Field(min_length=1, max_length=MAX_ROWS)
    mappings: list[RegionMapping] = Field(min_length=1, max_length=3)
    threshold: int = Field(default=100, ge=50, le=100, strict=True)
    combine: Literal['all', 'any'] = 'all'
    target_combine: Literal['all', 'any'] = 'any'
    match_mode: Literal['full', 'trim_suffix'] = 'full'
    preserve_short: bool = True
    suffixes: list[Literal['省', '市', '区', '县', '镇', '乡', '自治州', '自治县', '自治区']] = Field(default_factory=list, max_length=9)

    @model_validator(mode='after')
    def valid_table(self):
        if self.threshold % 10:
            raise ValueError('匹配精度请选择50%、60%、70%、80%、90%或100%')
        columns = [m.source_column for m in self.mappings]
        if len(set(columns)) != len(columns):
            raise ValueError('每个来源列只能配置一条映射，可在其中选择多个目标字段')
        for row in self.rows:
            if len(row) > MAX_COLUMNS or any(i >= len(row) for i in columns):
                raise ValueError('来源列超出数据列范围')
            if any(len(value) > 1000 for value in row):
                raise ValueError('单元格内容过长，请仅提交区域数据')
            if any(len(row[i]) > 255 for i in columns):
                raise ValueError('用于匹配的单元格最多255个字符')
        return self


def normalize(value):
    # 不删除内部空白或标点，避免把目标中分散的文字拼成连续字符串。
    return unicodedata.normalize('NFKC', str(value or '')).casefold().strip()


def prepare_source(value, payload):
    source = normalize(value)
    if payload.match_mode == 'trim_suffix' and not (payload.preserve_short and len(source) <= 2):
        # 自治区等优先作为完整后缀；只去一次，不递归剥离。
        for suffix in sorted(set(payload.suffixes), key=len, reverse=True):
            if source.endswith(suffix):
                return source[:-len(suffix)].strip()
    return source


class RegionIndex:
    def __init__(self, records):
        self.values = {field: defaultdict(set) for field in FIELDS}
        self.postings = {field: defaultdict(set) for field in FIELDS}
        self.cache = {}
        for record in records:
            for field in FIELDS:
                value = normalize(record[field])
                if value:
                    self.values[field][value].add(record['id'])
        for field in FIELDS:
            for value in self.values[field]:
                for char in set(value):
                    self.postings[field][char].add(value)

    def find_scores(self, source, field, threshold):
        # 兼容旧客户端 threshold 参数，但不再允许部分字符命中。
        key = (source, field)
        if key not in self.cache:
            candidates = [self.postings[field].get(char, set()) for char in set(source)]
            candidates.sort(key=len)
            values = set.intersection(*candidates) if candidates else set()
            found = {}
            for value in values:
                if source in value:
                    found.update({person_id: 100.0 for person_id in self.values[field][value]})
            self.cache[key] = found
        return self.cache[key]

    @staticmethod
    def combine_scores(groups, mode):
        if not groups:
            return {}
        ids = set.union(*(set(group) for group in groups)) if mode == 'any' else set.intersection(*(set(group) for group in groups))
        return {person_id: max(group.get(person_id, 0) for group in groups) if mode == 'any'
                else sum(group[person_id] for group in groups) / len(groups) for person_id in ids}

    def row_scores(self, values, payload):
        groups = []
        for mapping in payload.mappings:
            original = normalize(values[mapping.source_column])
            source = prepare_source(values[mapping.source_column], payload)
            if not source:
                # 原本非空但删成空字符串的条件不能变成“匹配所有”。
                if original:
                    groups.append({})
                continue
            targets = [self.find_scores(source, field, payload.threshold) for field in dict.fromkeys(mapping.target_fields)]
            groups.append(self.combine_scores(targets, payload.target_combine))
        return self.combine_scores(groups, payload.combine)

    def row(self, values, payload):
        return set(self.row_scores(values, payload))


def average_rate(scores):
    return round(sum(scores.values()) / len(scores), 2) if scores else None


def database_index(db):
    from resource_models import ResourcePerson
    from resource_service import _talent_query
    columns = [ResourcePerson.id, *[getattr(ResourcePerson, field) for field in FIELDS]]
    records = _talent_query(db).with_entities(*columns).all()
    return RegionIndex([dict(row._mapping) for row in records])


def match_summary(db, payload):
    index = database_index(db)
    unique_people = set()
    results = []
    for number, row in enumerate(payload.rows):
        matched = index.row_scores(row, payload)
        unique_people.update(matched)
        results.append({'index': number, 'total': len(matched), 'average_rate': average_rate(matched)})
    return {'items': results, 'unique_total': len(unique_people), 'total_rows': len(results)}


def match_page(db, payload, skip, limit):
    if len(payload.rows) != 1:
        raise HTTPException(422, '展开结果时一次只查询一行来源数据')
    scores = database_index(db).row_scores(payload.rows[0], payload)
    ids = set(scores)
    from resource_service import get_talents
    people = get_talents(db, person_ids=list(ids), skip=skip, limit=limit) if ids else []
    return {'items': people, 'total': len(ids), 'average_rate': average_rate(scores), 'rates': scores}


def column_name(index):
    result = ''
    while index >= 0:
        result = chr(65 + index % 26) + result
        index = index // 26 - 1
    return result


def parse_table(matrix, header):
    rows = []
    for source_row, values in enumerate(matrix, 1):
        values = ['' if value is None else str(value) for value in values]
        while values and not values[-1].strip():
            values.pop()
        if len(values) > MAX_COLUMNS:
            raise HTTPException(422, f'最多支持{MAX_COLUMNS}列，请只保留需要匹配的区域数据')
        if any(len(value) > 1000 for value in values):
            raise HTTPException(422, '单元格内容过长，请只上传区域数据')
        if values:
            rows.append({'source_row': source_row, 'values': values})
        if len(rows) > MAX_ROWS + int(header):
            raise HTTPException(422, f'一次最多匹配{MAX_ROWS}行，请拆分表格')
    headings = rows.pop(0)['values'] if header and rows else []
    width = max([len(headings), *[len(row['values']) for row in rows]], default=0)
    if not rows or not width:
        raise HTTPException(422, '没有可匹配的数据，请检查表头和工作表')
    for row in rows:
        row['values'] += [''] * (width - len(row['values']))
    return {'columns': [{'index': i, 'label': f'{column_name(i)}列 · {headings[i] if i < len(headings) and headings[i].strip() else column_name(i)}'} for i in range(width)], 'rows': rows}


def import_table(content, filename, sheet_name=None, header=True):
    suffix = Path(filename or '').suffix.lower()
    if len(content) > MAX_BYTES:
        raise HTTPException(413, '文件不能超过10MB')
    if suffix in {'.csv', '.tsv'}:
        try:
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = content.decode('gb18030')
            delimiter = '\t' if suffix == '.tsv' else ','
            result = parse_table(csv.reader(StringIO(text), delimiter=delimiter), header)
            return {**result, 'sheets': [], 'selected_sheet': ''}
        except UnicodeError:
            raise HTTPException(422, '无法识别文本编码，请另存为UTF-8 CSV')
        except csv.Error:
            raise HTTPException(422, 'CSV内容无法读取，请检查单元格长度与引号格式')
    if suffix not in {'.xlsx', '.xlsm'}:
        raise HTTPException(422, '支持.xlsx、.xlsm、.csv、.tsv；旧版.xls请先另存为.xlsx')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 64 * 1024 * 1024 or len(archive.infolist()) > 2000:
                raise HTTPException(413, '工作簿解压后过大，请仅保留区域数据')
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            selected = sheet_name or workbook.sheetnames[0]
            if selected not in workbook.sheetnames:
                raise HTTPException(422, '工作表不存在，请重新选择')
            sheet = workbook[selected]
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise HTTPException(422, f'工作表最多支持{MAX_COLUMNS}列，请仅保留区域数据')
            if sheet.max_row and sheet.max_row > 10001:
                raise HTTPException(422, '工作表行数过多，请复制需要的区域到新表后上传')
            result = parse_table(sheet.iter_rows(values_only=True), header)
            return {**result, 'sheets': workbook.sheetnames, 'selected_sheet': selected}
        finally:
            workbook.close()
    except HTTPException:
        raise
    except (BadZipFile, OSError, ValueError, KeyError, IndexError, TypeError):
        raise HTTPException(422, '工作簿无法读取，请确认文件未加密且格式正确')
