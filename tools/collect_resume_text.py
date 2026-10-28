"""只读采集共享目录文件原文；输出仅保存到仓库外目录。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, as_completed, wait
from datetime import datetime, timedelta, timezone
from pathlib import Path


TZ = timezone(timedelta(hours=8))
SUPPORTED = {".pdf", ".docx", ".doc", ".txt", ".rtf", ".xlsx", ".xls", ".pptx", ".odt"}
NO_OCR = {".jpg", ".jpeg", ".png", ".webp", ".jfif", ".bmp", ".tif", ".tiff"}
ARCHIVES = {".zip", ".rar", ".7z"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
DOC_TEMP_LIMIT_BYTES = 1024 * 1024 * 1024
STALE_DOC_TEMP_SECONDS = 180


class DocTempLimitExceeded(RuntimeError):
    pass


def terminate_process_tree(process: subprocess.Popen) -> None:
    if os.name == "nt":
        subprocess.run(
            ["taskkill.exe", "/PID", str(process.pid), "/T", "/F"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=10, check=False,
        )
    elif process.poll() is None:
        process.kill()
    try:
        process.wait(timeout=5)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=5)


def directory_size(path: str) -> int:
    total = 0
    for directory, _, files in os.walk(path):
        for name in files:
            try:
                total += os.stat(os.path.join(directory, name)).st_size
            except OSError:
                continue
    return total


def cleanup_stale_doc_temps() -> None:
    temp_root = os.environ.get("RESUME_TEMP_DIR")
    if not temp_root:
        return
    root = Path(temp_root).resolve()
    if not root.is_dir():
        return
    threshold = time.time() - STALE_DOC_TEMP_SECONDS
    with os.scandir(root) as entries:
        for entry in entries:
            if not re.fullmatch(r"resume-doc-[a-z0-9_]{8}", entry.name):
                continue
            try:
                if not entry.is_dir(follow_symlinks=False) or entry.stat(follow_symlinks=False).st_mtime > threshold:
                    continue
                target = Path(entry.path).resolve()
                if target.parent != root:
                    continue
                shutil.rmtree(target)
            except OSError:
                # 刚释放的 Office 文件句柄可能仍被 Windows 占用，下次再试。
                continue


def now_iso() -> str:
    return datetime.now(TZ).isoformat(timespec="seconds")


def file_time(timestamp: float) -> str:
    return datetime.fromtimestamp(timestamp, TZ).isoformat(timespec="seconds")


def file_id(path: str) -> str:
    return hashlib.sha256(os.path.normcase(os.path.abspath(path)).encode("utf-8")).hexdigest()


def sha256_file(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def segments_from_pdf(path: str) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        if not reader.decrypt(""):
            raise ValueError("encrypted_pdf")
    return [
        {"location": f"第 {index} 页", "text": page.extract_text(extraction_mode="plain") or ""}
        for index, page in enumerate(reader.pages, 1)
    ]


def segments_from_docx(path: str) -> list[dict]:
    result = []
    with zipfile.ZipFile(path) as archive:
        members = [name for name in archive.namelist() if name == "word/document.xml"]
        members += sorted(name for name in archive.namelist() if name.startswith("word/header") or name.startswith("word/footer"))
        for name in members:
            root = ET.fromstring(archive.read(name))
            for index, paragraph in enumerate(root.iter(W + "p"), 1):
                value = "".join(node.text or "" for node in paragraph.iter(W + "t"))
                if value.strip():
                    result.append({"location": f"{name} 第 {index} 段", "text": value})
    return result


def segments_from_doc(path: str) -> list[dict]:
    # 每个子进程使用独立 LibreOffice 配置，避免与服务器桌面 Office 会话争锁。
    office = Path(r"C:\Program Files\LibreOffice\program\soffice.com")
    if not office.is_file():
        raise RuntimeError("libreoffice_missing")
    temp_root = os.environ.get("RESUME_TEMP_DIR")
    if not temp_root or not Path(temp_root).is_dir():
        raise RuntimeError("resume_temp_dir_missing")
    with tempfile.TemporaryDirectory(prefix="resume-doc-", dir=temp_root) as directory:
        profile = (Path(directory) / "profile").as_uri()
        process = subprocess.Popen(
            [str(office), f"-env:UserInstallation={profile}", "--headless", "--convert-to", "txt:Text (encoded):UTF8", "--outdir", directory, path],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        started = time.monotonic()
        try:
            while process.poll() is None:
                if directory_size(directory) > DOC_TEMP_LIMIT_BYTES:
                    raise DocTempLimitExceeded("doc_temp_limit")
                if time.monotonic() - started > 100:
                    raise subprocess.TimeoutExpired(process.args, 100)
                try:
                    process.wait(timeout=1)
                except subprocess.TimeoutExpired:
                    pass
        except (DocTempLimitExceeded, subprocess.TimeoutExpired):
            terminate_process_tree(process)
            raise
        if directory_size(directory) > DOC_TEMP_LIMIT_BYTES:
            raise DocTempLimitExceeded("doc_temp_limit")
        converted = Path(directory) / (Path(path).stem + ".txt")
        if process.returncode or not converted.is_file():
            raise RuntimeError("doc_conversion_failed")
        return [{"location": "正文", "text": converted.read_text(encoding="utf-8-sig")}]


def segments_from_txt(path: str) -> list[dict]:
    data = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return [{"location": "全文", "text": data.decode(encoding)}]
        except UnicodeDecodeError:
            continue
    raise ValueError("text_encoding_unknown")


def segments_from_rtf(path: str) -> list[dict]:
    from striprtf.striprtf import rtf_to_text

    data = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            return [{"location": "全文", "text": rtf_to_text(data.decode(encoding))}]
        except UnicodeDecodeError:
            continue
    raise ValueError("rtf_encoding_unknown")


def segments_from_xlsx(path: str) -> list[dict]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        result = []
        for sheet in book:
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                values = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                if values:
                    result.append({"location": f"{sheet.title} 第 {row_index} 行", "text": "\t".join(values)})
        return result
    finally:
        book.close()


def segments_from_xls(path: str) -> list[dict]:
    import xlrd

    book = xlrd.open_workbook(path, on_demand=True)
    try:
        result = []
        for sheet in book.sheets():
            for row_index in range(sheet.nrows):
                values = [f"列 {col_index + 1}={sheet.cell_value(row_index, col_index)}" for col_index in range(sheet.ncols) if sheet.cell_value(row_index, col_index) != ""]
                if values:
                    result.append({"location": f"{sheet.name} 第 {row_index + 1} 行", "text": "\t".join(values)})
        return result
    finally:
        book.release_resources()


def segments_from_pptx(path: str) -> list[dict]:
    result = []
    with zipfile.ZipFile(path) as archive:
        slides = sorted((name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")), key=lambda name: int(name.rsplit("slide", 1)[1].split(".")[0]))
        for index, name in enumerate(slides, 1):
            root = ET.fromstring(archive.read(name))
            lines = [node.text or "" for node in root.iter(A + "t")]
            result.append({"location": f"第 {index} 页", "text": "\n".join(lines)})
    return result


def segments_from_odt(path: str) -> list[dict]:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("content.xml"))
    text_ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    return [{"location": f"第 {i} 段", "text": "".join(node.itertext())} for i, node in enumerate(root.iter(text_ns + "p"), 1)]


EXTRACTORS = {
    ".pdf": segments_from_pdf, ".docx": segments_from_docx, ".doc": segments_from_doc,
    ".txt": segments_from_txt, ".rtf": segments_from_rtf, ".xlsx": segments_from_xlsx,
    ".xls": segments_from_xls, ".pptx": segments_from_pptx, ".odt": segments_from_odt,
}


def worker(path: str, result_path: str, text_path: str, segments_path: str) -> int:
    try:
        digest = sha256_file(path)
        extension = Path(path).suffix.lower()
        if Path(path).name.startswith("~$"):
            status, reason, segments = "skipped", "office_lock_file", []
        elif extension in NO_OCR:
            status, reason, segments = "skipped", "image_requires_ocr", []
        elif extension in ARCHIVES:
            status, reason, segments = "skipped", "archive_not_opened", []
        elif extension not in SUPPORTED:
            status, reason, segments = "skipped", "unsupported_format", []
        else:
            segments = EXTRACTORS[extension](path)
            status = "success" if any(item["text"].strip() for item in segments) else "no_text"
            reason = "" if status == "success" else "empty_or_scanned"
        if status == "success":
            Path(text_path).write_text("\n\n".join(item["text"] for item in segments if item["text"].strip()), encoding="utf-8")
            Path(segments_path).write_text(json.dumps(segments, ensure_ascii=False), encoding="utf-8")
        result = {"status": status, "reason": reason, "sha256": digest, "characters": sum(len(item["text"]) for item in segments), "segments": len(segments)}
    except Exception as exc:
        # 不将异常消息写入报告；第三方库异常可能包含简历正文或路径。
        result = {"status": "failed", "reason": type(exc).__name__, "sha256": "", "characters": 0, "segments": 0}
    Path(result_path).write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    return 0


def ensure_output_safe(output: Path, roots: list[str]) -> None:
    resolved = output.resolve()
    for root in roots:
        if os.path.normcase(str(resolved)).startswith(os.path.normcase(os.path.abspath(root)).rstrip("\\/") + os.sep):
            raise ValueError("输出目录不能位于采集目录内部")
    repository = Path(__file__).resolve().parents[1]
    if resolved == repository or repository in resolved.parents:
        raise ValueError("输出目录必须位于项目仓库之外")


def database(connection: sqlite3.Connection) -> None:
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("""CREATE TABLE IF NOT EXISTS files (
        path TEXT PRIMARY KEY, root TEXT NOT NULL, relative_path TEXT NOT NULL,
        extension TEXT NOT NULL, size INTEGER NOT NULL, mtime_ns INTEGER NOT NULL,
        modified_at TEXT NOT NULL, sha256 TEXT NOT NULL, status TEXT NOT NULL,
        reason TEXT NOT NULL, characters INTEGER NOT NULL, segments INTEGER NOT NULL,
        processed_at TEXT NOT NULL, file_id TEXT NOT NULL
    )""")
    connection.commit()


def iter_files(root: str, errors: list[dict]):
    def onerror(exc: OSError):
        errors.append({"root": root, "path": exc.filename or root, "reason": type(exc).__name__})

    if os.path.isfile(root):
        yield root, os.path.basename(root)
        return
    for directory, folders, files in os.walk(root, onerror=onerror):
        folders.sort(key=str.casefold)
        for name in sorted(files, key=str.casefold):
            path = os.path.join(directory, name)
            yield path, os.path.relpath(path, root)


def process_one(path: str, root: str, relative: str, stat: os.stat_result, output: Path, timeout: int) -> dict:
    identifier = file_id(path)
    with tempfile.TemporaryDirectory(prefix="resume-worker-", dir=output / "tmp") as temp:
        temp_path = Path(temp)
        result_path, text_path, segments_path = (temp_path / name for name in ("result.json", "text.txt", "segments.json"))
        try:
            process = subprocess.Popen(
                [sys.executable, str(Path(__file__).resolve()), "--worker", path, str(result_path), str(text_path), str(segments_path)],
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
            try:
                process.wait(timeout=timeout)
            except subprocess.TimeoutExpired:
                terminate_process_tree(process)
                raise
            if process.returncode or not result_path.exists():
                result = {"status": "failed", "reason": "worker_exit", "sha256": "", "characters": 0, "segments": 0}
            else:
                result = json.loads(result_path.read_text(encoding="utf-8"))
        except subprocess.TimeoutExpired:
            result = {"status": "failed", "reason": "timeout", "sha256": "", "characters": 0, "segments": 0}
        if result["status"] == "success":
            os.replace(text_path, output / "texts" / f"{identifier}.txt")
            os.replace(segments_path, output / "segments" / f"{identifier}.json")
        else:
            (output / "texts" / f"{identifier}.txt").unlink(missing_ok=True)
            (output / "segments" / f"{identifier}.json").unlink(missing_ok=True)
    return {
        "path": path, "root": root, "relative_path": relative, "extension": Path(path).suffix.lower(),
        "size": stat.st_size, "mtime_ns": stat.st_mtime_ns, "modified_at": file_time(stat.st_mtime),
        "sha256": result["sha256"], "status": result["status"], "reason": result["reason"],
        "characters": result["characters"], "segments": result["segments"],
        "processed_at": now_iso(), "file_id": identifier,
    }


def export(connection: sqlite3.Connection, output: Path, roots: list[str], scan_errors: list[dict], current_paths: set[str]) -> dict:
    columns = ["path", "root", "relative_path", "extension", "size", "mtime_ns", "modified_at", "sha256", "status", "reason", "characters", "segments", "processed_at", "file_id"]
    query = "SELECT " + ",".join(columns) + " FROM files ORDER BY root,relative_path"
    manifest = output / "manifest.jsonl"
    pending = output / "pending.csv"
    with manifest.open("w", encoding="utf-8") as target, pending.open("w", encoding="utf-8-sig", newline="") as pending_target:
        writer = csv.DictWriter(pending_target, fieldnames=["path", "status", "reason"])
        writer.writeheader()
        for values in connection.execute(query):
            item = dict(zip(columns, values))
            if item["path"] not in current_paths:
                continue
            item["text_file"] = f"texts/{item['file_id']}.txt" if item["status"] == "success" else None
            item["segments_file"] = f"segments/{item['file_id']}.json" if item["status"] == "success" else None
            target.write(json.dumps(item, ensure_ascii=False) + "\n")
            if item["status"] != "success":
                writer.writerow({key: item[key] for key in writer.fieldnames})
    counts = Counter()
    by_root = {root: Counter() for root in roots}
    for values in connection.execute(query):
        item = dict(zip(columns, values))
        if item["path"] in current_paths:
            counts[item["status"]] += 1
            by_root[item["root"]][item["status"]] += 1
    report = {"generated_at": now_iso(), "roots": roots, "total": len(current_paths), "counts": dict(counts), "by_root": {root: dict(value) for root, value in by_root.items()}, "scan_errors": scan_errors}
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def collect(roots: list[str], output: Path, workers: int, timeout: int) -> int:
    roots = [os.path.abspath(root) for root in roots]
    ensure_output_safe(output, roots)
    output.mkdir(parents=True, exist_ok=True)
    for name in ("texts", "segments", "tmp"):
        (output / name).mkdir(exist_ok=True)
    connection = sqlite3.connect(output / "state.sqlite3")
    database(connection)
    cleanup_stale_doc_temps()
    errors: list[dict] = []
    current_paths: set[str] = set()
    current_keys: set[str] = set()
    processed = 0
    queued = {}

    def save_completed(future) -> None:
        nonlocal processed
        path, root, relative, stat = queued.pop(future)
        try:
            item = future.result()
        except Exception as exc:
            item = {
                "path": path, "root": root, "relative_path": relative, "extension": Path(path).suffix.lower(),
                "size": stat.st_size, "mtime_ns": stat.st_mtime_ns, "modified_at": file_time(stat.st_mtime),
                "sha256": "", "status": "failed", "reason": type(exc).__name__, "characters": 0,
                "segments": 0, "processed_at": now_iso(), "file_id": file_id(path),
            }
        try:
            connection.execute("""INSERT INTO files VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(path) DO UPDATE SET root=excluded.root,relative_path=excluded.relative_path,
                extension=excluded.extension,size=excluded.size,mtime_ns=excluded.mtime_ns,
                modified_at=excluded.modified_at,sha256=excluded.sha256,status=excluded.status,
                reason=excluded.reason,characters=excluded.characters,segments=excluded.segments,
                processed_at=excluded.processed_at,file_id=excluded.file_id""", tuple(item.values()))
            connection.commit()
        finally:
            processed += 1
            if processed % 25 == 0:
                cleanup_stale_doc_temps()
            if processed % 100 == 0:
                print(f"已处理 {processed}；累计发现 {len(current_paths)} 个文件", flush=True)

    with ThreadPoolExecutor(max_workers=workers) as executor:
        for root in roots:
            if not os.path.exists(root):
                errors.append({"root": root, "path": root, "reason": "root_unavailable"})
                continue
            for path, relative in iter_files(root, errors):
                key = os.path.normcase(os.path.abspath(path))
                if key in current_keys:
                    continue
                try:
                    stat = os.stat(path)
                except OSError as exc:
                    errors.append({"root": root, "path": path, "reason": type(exc).__name__})
                    continue
                current_keys.add(key)
                current_paths.add(path)
                old = connection.execute("SELECT size,mtime_ns,status,file_id FROM files WHERE path=?", (path,)).fetchone()
                if old and old[:2] == (stat.st_size, stat.st_mtime_ns) and old[2] != "failed" and (old[2] != "success" or (output / "texts" / f"{old[3]}.txt").exists()):
                    continue
                future = executor.submit(process_one, path, root, relative, stat, output, timeout)
                queued[future] = (path, root, relative, stat)
                if len(queued) >= workers * 4:
                    done, _ = wait(queued, return_when=FIRST_COMPLETED)
                    for completed in done:
                        save_completed(completed)
        for completed in as_completed(queued):
            save_completed(completed)
    if errors:
        # 扫描不完整时仍保存已发现结果，但返回非零，避免误报完整采集。
        print(f"目录扫描存在 {len(errors)} 项错误；详见 report.json", file=sys.stderr)
    report = export(connection, output, roots, errors, current_paths)
    connection.close()
    print(json.dumps(report, ensure_ascii=False), flush=True)
    return 1 if errors or report["counts"].get("failed", 0) else 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", action="append", default=[], help="采集根路径；可重复指定")
    parser.add_argument("--output", type=Path, help="仓库外输出目录")
    parser.add_argument("--workers", type=int, default=2)
    parser.add_argument("--timeout", type=int, default=120, help="每文件处理超时（秒）")
    parser.add_argument("--worker", nargs=4, metavar=("SOURCE", "RESULT", "TEXT", "SEGMENTS"), help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.worker:
        return worker(*args.worker)
    if not args.root or args.output is None:
        parser.error("必须提供 --root 和 --output")
    if not 1 <= args.workers <= 4 or args.timeout < 5:
        parser.error("--workers 须为 1-4，--timeout 不小于 5 秒")
    return collect(args.root, args.output, args.workers, args.timeout)


if __name__ == "__main__":
    raise SystemExit(main())
