"""口译、标注、招聘项目导出：沿用列表查询，按批次写出已有业务数据。"""

from __future__ import annotations

import json
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import selectinload

from project_export_excel import ExportColumn, _append_row, _make_sheet

EXPORT_BATCH_SIZE = 500
EXPORT_MAX_ROWS_PER_SHEET = 50_000
MODULE_LABELS = {"interpretation": "口译项目", "annotation": "标注项目", "recruitment": "招聘项目"}
COMMON_TIME_FIELDS = {
    "customer_consultation_time": "客户咨询时间",
    "customer_confirmation_time": "客户确认时间",
    "created_at": "创建时间",
}
TIME_FIELDS = {
    "interpretation": {**COMMON_TIME_FIELDS, "scheduled_date": "预定日期"},
    "annotation": {**COMMON_TIME_FIELDS, "task_dispatched_at": "任务派发时间", "task_submitted_at": "任务提交时间"},
    "recruitment": {**COMMON_TIME_FIELDS, "target_onboard_date": "目标入职日期"},
}


class ProjectExportEmptyError(ValueError):
    pass


class ProjectExportLimitError(ValueError):
    pass


def read(item, path, default=None):
    for part in path.split("."):
        item = item.get(part) if isinstance(item, dict) else getattr(item, part, None)
        if item is None:
            return default
    return item


def text_value(value):
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (list, tuple)):
        return "；".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


# 显式字段白名单；不读取账号密码、内部幂等键或结算推算值。
COMMON_FIELDS = [
    ("order_no", "订单号", "identifier"), ("project_name", "项目名称"),
    ("project_status", "项目状态"), ("client_full_name", "客户全称"),
    ("client_short_name", "客户简称"), ("client_code", "客户编号", "identifier"),
    ("contact_name", "联系人"), ("customer_order_no", "客户单号/项目标识", "identifier"),
    ("customer_consultation_time", "客户咨询时间", "datetime"),
    ("customer_confirmation_time", "客户确认时间", "datetime"),
    ("role_assignments", "项目职责"), ("quotation_path", "报价单路径"),
    ("contract_path", "合同路径"), ("email_subject_preview", "邮件主题"),
    ("created_at", "创建时间", "datetime"), ("updated_at", "更新时间", "datetime"),
]
PROJECT_FIELDS = {
    "interpretation": [
        ("project_types", "项目类型"), ("task_description", "任务描述"),
        ("locations", "地点"), ("customer_budget", "客户预算原文"),
        ("required_interpreter_count", "所需译员人数", "integer"),
        ("required_interpreter_gender", "所需译员性别"),
        ("required_interpretation_level", "所需口译级别"),
        ("interpreter_special_requirements", "译员特殊要求"),
        ("interpreter_height_requirement", "译员身高要求"),
        ("interpreter_appearance_requirement", "译员外貌要求"),
        ("interpreter_dress_requirement", "译员着装要求"),
        ("interpretation_domain", "口译领域"), ("interpretation_content", "口译内容"),
        ("client_domain", "客户领域"), ("current_client_manager", "客户经理"),
        ("manager_contact", "客户经理联系方式"), ("file_path", "项目文件路径"),
        ("client_rating", "客户对信实评价"), ("client_rating_note", "客户评价说明"),
        ("social_post_request", "发帖需求"), ("resource_request", "资源需求"), ("remarks", "备注"),
    ],
    "annotation": [
        ("parent_order_no", "母订单号", "identifier"), ("parent_project_name", "母项目名称"),
        ("project_types", "项目类型"), ("task_description", "任务描述"),
        ("priority", "优先次序"), ("language_region", "语言地区"),
        ("potential_demand", "潜在需求量原文"), ("client_manager_name", "客户经理"),
        ("project_path", "项目路径"), ("task_dispatched_at", "任务派发时间", "datetime"),
        ("task_submitted_at", "任务提交时间", "datetime"),
        ("status_effective_on", "状态发生时间", "datetime"), ("custom_values", "自定义字段"),
    ],
    "recruitment": [
        ("client_domain", "客户领域"), ("candidate_count", "候选人数量", "integer"),
        ("position_title", "招聘职位"), ("job_description", "岗位描述"),
        ("headcount_min", "招聘人数下限", "integer"), ("headcount_max", "招聘人数上限", "integer"),
        ("target_onboard_type", "目标入职方式"), ("target_onboard_date", "目标入职日期", "date"),
        ("employment_start", "用工开始日期", "date"), ("employment_end", "用工结束日期", "date"),
        ("work_location", "工作地点"), ("client_manager_name", "客户经理"),
        ("service_fee_type", "服务费类型"), ("service_fee_currency", "服务费币种"),
        ("service_fee_amount", "固定服务费", "money"), ("service_fee_rate", "服务费率（%）", "decimal"),
        ("service_fee_multiplier", "服务费倍数", "decimal"), ("service_fee_note", "服务费说明"),
        ("project_path", "项目路径"), ("social_post_request", "发帖需求"),
        ("resource_request", "资源需求"), ("remarks", "备注"),
    ],
}
TIME_RANGE_FIELDS = [
    ("sequence_no", "安排序号", "integer"), ("scheduled_start", "预定开始时间", "datetime"),
    ("scheduled_end", "预定结束时间", "datetime"), ("actual_start", "实际开始时间", "datetime"),
    ("actual_end", "实际结束时间", "datetime"),
]
INTERPRETER_FIELDS = [
    ("sequence_no", "安排序号", "integer"), ("translator_name", "译员姓名"),
    ("translator_code", "译员编号", "identifier"), ("translator_gender", "译员性别"),
    ("translator_languages", "译员语言"), ("translator_interpretation_level", "口译级别"),
    ("translator_height", "译员身高"), ("translator_appearance", "译员外貌"),
    ("translator_direction", "译员方向"), ("translator_translation_type", "译员翻译类型"),
    ("translator_resume_path", "译员简历路径"), ("customer_rating", "客户对译员评价"),
    ("evaluation_note", "译员评价说明"),
]
LANGUAGE_FIELDS = {
    "interpretation": [("sequence_no", "语言序号", "integer"), ("display", "语言方向"), ("required_count", "所需人数", "integer")],
    "annotation": [("sequence_no", "语言序号", "integer"), ("display", "语言项")],
    "recruitment": [("direction_type", "语言需求类型"), ("label", "语言需求")],
}
PRICE_FIELDS = [
    ("sequence_no", "报价序号", "integer"), ("project_type", "报价项目类型"),
    ("language_display", "报价语言"), ("amount", "客户单价", "decimal"),
    ("currency", "报价币种"), ("unit", "客户计价单位"), ("remarks", "报价备注"),
]
ASSIGNEE_FIELDS = [
    ("sequence_no", "安排序号", "integer"), ("person_name", "人员姓名"),
    ("resource_code", "资源编号", "identifier"), ("assignment_role", "安排角色"),
    ("language_item.display", "安排语种"), ("assignment_status", "安排状态"),
    ("audio_duration_value", "已记录音频工作量", "decimal"), ("audio_duration_unit", "工作量单位"),
    ("rate.amount", "人员单价", "decimal"), ("rate.currency", "人员币种"), ("rate.unit", "人员计价单位"),
    ("rate.remarks", "单价备注"), ("quality_score", "质量评分"), ("evaluation_note", "人员评价说明"),
    ("custom_values", "人员自定义字段"), ("created_at", "安排创建时间", "datetime"),
    ("updated_at", "安排更新时间", "datetime"),
]
CANDIDATE_FIELDS = [
    ("id", "候选人记录ID", "identifier"), ("candidate_name", "候选人姓名"),
    ("contact_info", "联系方式"), ("resume_path", "简历路径"), ("resume_source_label", "简历来源"),
    ("stage", "候选人阶段"), ("recommended_at", "推荐时间", "datetime"),
    ("interview_at", "面试时间", "datetime"), ("offer_at", "Offer时间", "datetime"),
    ("planned_onboard_date", "计划入职日期", "date"), ("actual_onboard_date", "实际入职日期", "date"),
    ("first_interview_date", "一面日期", "date"), ("first_interview_details", "一面详情"),
    ("second_interview_date", "二面日期", "date"), ("second_interview_details", "二面详情"),
    ("owner_name", "负责人"), ("next_follow_up_at", "下次跟进时间", "datetime"),
    ("remarks", "候选人备注"), ("created_at", "候选人创建时间", "datetime"),
    ("updated_at", "候选人更新时间", "datetime"),
]
PROGRESS_FIELDS = [
    ("from_status", "原项目状态"), ("to_status", "变更后项目状态"), ("note", "进度说明"),
    ("operator_name", "操作人"), ("is_system", "系统记录"), ("occurred_at", "发生时间", "datetime"),
]
COMMUNICATION_FIELDS = [
    ("sequence_no", "沟通序号", "integer"), ("communication_date", "沟通日期", "date"),
    ("details", "沟通详情"), ("created_at", "沟通创建时间", "datetime"), ("updated_at", "沟通更新时间", "datetime"),
]
INTERVIEW_FIELDS = [
    ("round_no", "面试轮次", "integer"), ("interview_date", "面试日期", "date"),
    ("details", "面试详情"), ("created_at", "面试创建时间", "datetime"), ("updated_at", "面试更新时间", "datetime"),
]


def enum_labels(module):
    from interpretation_schemas import PROJECT_TYPE_LABELS
    from annotation_schemas import ANNOTATION_PROJECT_TYPE_LABELS, ANNOTATION_PROJECT_STATUS_LABELS, ANNOTATION_PROJECT_PRIORITY_LABELS
    from recruitment_schemas import PROJECT_STATUSES, CANDIDATE_STAGES
    statuses = {
        "initial_follow_up": "初步跟进中", "deal_pending_execution": "已成交待执行",
        "in_progress": "进行中", "cancelled": "已取消", "partially_cancelled": "已部分取消",
        "ended": "已结束", "settled": "已结款",
    } if module == "interpretation" else ANNOTATION_PROJECT_STATUS_LABELS if module == "annotation" else PROJECT_STATUSES
    return {
        "project_status": statuses, "from_status": statuses, "to_status": statuses,
        "project_types": PROJECT_TYPE_LABELS if module == "interpretation" else ANNOTATION_PROJECT_TYPE_LABELS,
        "project_type": ANNOTATION_PROJECT_TYPE_LABELS, "priority": ANNOTATION_PROJECT_PRIORITY_LABELS,
        "stage": CANDIDATE_STAGES, "assignment_role": {"annotator": "标注员", "quality_inspector": "质检员"},
        "assignment_status": {"assigned": "已安排", "in_progress": "进行中", "completed": "已完成", "cancelled": "已取消"},
        "direction_type": {"single": "单语种", "translation": "翻译方向"},
        "target_onboard_type": {"date": "指定日期", "anytime": "随时"},
        "service_fee_type": {"fixed": "固定费用", "monthly_salary_multiple": "月薪倍数", "annual_salary_rate": "年薪比例", "other": "其他"},
        "audio_duration_unit": {"second": "秒", "minute": "分钟", "hour": "小时"},
        "rate.unit": {"item": "条", "second": "秒", "minute": "分钟", "hour": "小时"},
        "required_interpreter_gender": {"male": "男", "female": "女", "any": "不限"},
        "required_interpretation_level": {"junior": "初级", "intermediate": "中级", "senior": "高级"},
        "client_rating": {"very_satisfied": "非常满意", "satisfied": "满意", "basically_satisfied": "基本满意", "dissatisfied": "不满意", "very_dissatisfied": "非常不满意"},
    }


def values(item, fields, labels, custom_labels=None):
    result = {}
    for field, label, *_kind in fields:
        value = read(item, field)
        mapping = labels.get("client_rating" if field == "customer_rating" else field)
        if mapping:
            value = [mapping.get(str(v), v) for v in value] if isinstance(value, list) else mapping.get(str(value), value)
        if field == "role_assignments":
            value = "；".join(f'{read(role, "role_name") or read(role, "role_code")}：{read(role, "assignee_name") or "角色池"}' for role in (value or []))
        if field == "custom_values" and isinstance(value, dict):
            value = {str((custom_labels or {}).get(str(key), key)): text_value(v) for key, v in value.items()}
        result[label] = value if _kind and _kind[0] in {"date", "datetime", "integer", "decimal", "money"} else text_value(value)
    return result


def columns(fields):
    return [ExportColumn(label, lambda row, key=label: row.get(key), kind[0] if kind else "text", 24)
            for _, label, *kind in fields]


@dataclass(frozen=True)
class SheetSpec:
    title: str
    fields: list


def project_fields(module):
    # 招聘现有模型将客户全称命名为 client_name，其他模块为 client_full_name。
    common = [("client_name", *field[1:]) if module == "recruitment" and field[0] == "client_full_name" else field for field in COMMON_FIELDS]
    return common + PROJECT_FIELDS[module]


def sheet_specs(module, mode):
    fields = project_fields(module)
    if mode in {"reconciliation", "personnel"}:
        detail = PRICE_FIELDS if module == "annotation" and mode == "reconciliation" else ASSIGNEE_FIELDS if module == "annotation" else INTERPRETER_FIELDS if mode == "personnel" else []
        reconciliation_fields = fields + detail
        return [SheetSpec("对账单", reconciliation_fields), SheetSpec("待补数据", reconciliation_fields + [("missing", "待补原因")])]
    if module == "interpretation":
        return [SheetSpec("项目", fields), SheetSpec("时间安排", fields + TIME_RANGE_FIELDS),
                SheetSpec("语言方向", fields + LANGUAGE_FIELDS[module]), SheetSpec("译员安排", fields + INTERPRETER_FIELDS)]
    if module == "annotation":
        return [SheetSpec("母订单", fields), SheetSpec("子订单", fields),
                SheetSpec("语言项", fields + LANGUAGE_FIELDS[module]), SheetSpec("客户单价", fields + PRICE_FIELDS),
                SheetSpec("人员安排", fields + ASSIGNEE_FIELDS)]
    candidate_fields = fields + CANDIDATE_FIELDS
    specs = [SheetSpec("候选人", candidate_fields), SheetSpec("候选人沟通", candidate_fields + COMMUNICATION_FIELDS),
             SheetSpec("面试记录", candidate_fields + INTERVIEW_FIELDS)]
    if mode == "projects":
        specs = [SheetSpec("项目", fields), SheetSpec("语言需求", fields + LANGUAGE_FIELDS[module]),
                 SheetSpec("项目进度", fields + PROGRESS_FIELDS)] + specs
    return specs


def project_rows(module, mode, project, custom_labels=None):
    labels = enum_labels(module)
    base = values(project, project_fields(module), labels, custom_labels)
    if mode in {"reconciliation", "personnel"}:
        if mode == "personnel":
            if read(project, "project_status") == "cancelled":
                return
            relation, fields = ("assignees", ASSIGNEE_FIELDS) if module == "annotation" else ("interpreter_assignments", INTERPRETER_FIELDS)
            details = [item for item in read(project, relation, []) if read(item, "assignment_status") != "cancelled"]
        elif module == "annotation":
            details, fields = read(project, "price_items", []) or [None], PRICE_FIELDS
        else:
            details, fields = [None], []
        for detail in details:
            row = {**base, **values(detail, fields, labels, custom_labels)}
            missing = ["账单月份", "已确认结算金额", "含税/不含税结算依据"]
            if not base.get("客户全称"):
                missing.append("客户全称")
            if not base.get("客户编号"):
                missing.append("客户编号")
            if module == "annotation":
                if detail is None:
                    missing.append("客户单价")
                elif mode == "personnel" and read(detail, "rate.amount") is None:
                    missing.append("人员单价")
                if detail is not None and not read(detail, "rate.currency" if mode == "personnel" else "currency"):
                    missing.append("币种")
                missing.append("已确认计费数量")
            row["待补原因"] = "；".join(missing)
            yield "待补数据", row
        return
    if mode == "projects":
        title = ("子订单" if read(project, "parent_project_id") else "母订单") if module == "annotation" else "项目"
        yield title, base
    relations = {
        "interpretation": [("时间安排", "time_ranges", TIME_RANGE_FIELDS), ("语言方向", "language_directions", LANGUAGE_FIELDS[module]), ("译员安排", "interpreter_assignments", INTERPRETER_FIELDS)],
        "annotation": [("语言项", "language_items", LANGUAGE_FIELDS[module]), ("客户单价", "price_items", PRICE_FIELDS), ("人员安排", "assignees", ASSIGNEE_FIELDS)],
        "recruitment": [("语言需求", "language_directions", LANGUAGE_FIELDS[module]), ("项目进度", "progress_records", PROGRESS_FIELDS)],
    }[module]
    if mode == "projects":
        for title, relation, fields in relations:
            for item in read(project, relation, []):
                yield title, {**base, **values(item, fields, labels, custom_labels)}
    if module == "recruitment":
        for candidate in read(project, "candidates", []):
            row = {**base, **values(candidate, CANDIDATE_FIELDS, labels)}
            yield "候选人", row
            for title, relation, fields in [("候选人沟通", "communications", COMMUNICATION_FIELDS), ("面试记录", "interviews", INTERVIEW_FIELDS)]:
                for item in read(candidate, relation, []):
                    yield title, {**row, **values(item, fields, labels)}


def projects_to_xlsx(module: str, mode: str, batches: Iterable, *, custom_labels=None, max_rows_per_sheet=EXPORT_MAX_ROWS_PER_SHEET):
    specs = sheet_specs(module, mode)
    workbook = Workbook(write_only=True)
    sheets = {spec.title: (_make_sheet(workbook, spec.title, columns(spec.fields)), columns(spec.fields)) for spec in specs}
    counts = dict.fromkeys(sheets, 0)
    seen = set()
    try:
        for batch in batches:
            for project in batch:
                key = read(project, "id")
                if key is not None:
                    if str(key) in seen:
                        continue
                    seen.add(str(key))
                for title, row in project_rows(module, mode, project, custom_labels):
                    counts[title] += 1
                    if counts[title] > max_rows_per_sheet:
                        raise ProjectExportLimitError(f"{title}超过 {max_rows_per_sheet} 行，请缩小导出范围")
                    sheet, cols = sheets[title]
                    _append_row(sheet, cols, row)
        if not any(counts.values()):
            raise ProjectExportEmptyError("所选范围内没有可导出的数据")
    except Exception:
        # 提前失败时也完整关闭 write_only writer 的临时资源。
        try:
            workbook.save(BytesIO())
        except Exception:
            pass
        raise
    for title, (sheet, cols) in sheets.items():
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{counts[title] + 1}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def iter_project_batches(db, module, filters):
    from interpretation_service import get_interpretation_projects
    from annotation_service import get_annotation_projects, _project_options
    from annotation_models import AnnotationProject, AnnotationProjectAssignee, AnnotationProjectLanguageItem
    from recruitment_service import get_recruitment_projects, _project_detail_options
    from recruitment_models import RecruitmentProject, RecruitmentCandidate

    loaders = {"interpretation": get_interpretation_projects, "annotation": get_annotation_projects, "recruitment": get_recruitment_projects}
    extra = ()
    if module == "annotation":
        extra = (selectinload(AnnotationProject.assignees).selectinload(AnnotationProjectAssignee.rate),
                 selectinload(AnnotationProject.assignees).joinedload(AnnotationProjectAssignee.language_item).joinedload(AnnotationProjectLanguageItem.source_language),
                 selectinload(AnnotationProject.assignees).joinedload(AnnotationProjectAssignee.language_item).joinedload(AnnotationProjectLanguageItem.target_language))
    elif module == "recruitment":
        extra = (*_project_detail_options(), selectinload(RecruitmentProject.candidates).selectinload(RecruitmentCandidate.interviews))
    skip = 0
    while True:
        rows = loaders[module](db, skip=skip, limit=EXPORT_BATCH_SIZE, extra_options=extra, **filters)
        if not rows:
            break
        yield rows
        if module == "annotation" and filters.get("order_scope", "parent") == "parent":
            parent_ids = [item.id for item in rows if not item.parent_project_id]
            # 每个父批次的子单也分页读取，避免大型母订单一次装载所有子单。
            child_skip = 0
            while parent_ids:
                children = (db.query(AnnotationProject).options(*_project_options(), *extra)
                            .filter(AnnotationProject.parent_project_id.in_(parent_ids))
                            .order_by(AnnotationProject.order_no.desc(), AnnotationProject.id.desc())
                            .offset(child_skip).limit(EXPORT_BATCH_SIZE).all())
                if not children:
                    break
                yield children
                child_skip += len(children)
        skip += len(rows)
        if len(rows) < EXPORT_BATCH_SIZE:
            break


def create_project_export(db, module, mode, filters):
    custom_labels = {}
    if module == "annotation":
        from annotation_ops_models import AnnotationCustomFieldDefinition
        definitions = db.query(AnnotationCustomFieldDefinition).filter(
            AnnotationCustomFieldDefinition.table_code.in_(["project", "assignment"]),
        ).all()
        # 自定义值使用字段 ID；历史停用定义也保留可读名称。
        custom_labels = {str(item.id): item.field_label for item in definitions}
    return projects_to_xlsx(module, mode, iter_project_batches(db, module, filters), custom_labels=custom_labels)
