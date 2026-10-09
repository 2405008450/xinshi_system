"""项目共用的 Excel 类型、样式和文本安全处理。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Callable, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_ILLEGAL_EXCEL_TEXT = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_FORMULA_PREFIXES = ("=", "+", "-", "@")


@dataclass(frozen=True)
class ExportColumn:
    label: str
    getter: Callable[[Any], Any]
    kind: str = "text"
    width: int = 18


def _safe_text(value: Any) -> str | None:
    if value is None:
        return None
    text = _ILLEGAL_EXCEL_TEXT.sub("", str(value))[:32767]
    if text.startswith(_FORMULA_PREFIXES):
        text = f"'{text}"
    return text


def _percent_value(value: Any) -> float | str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value) / 100
    normalized = str(value).strip()
    try:
        return float(normalized.rstrip("%")) / 100
    except ValueError:
        return _safe_text(normalized)


def _cell_value(value: Any, kind: str) -> Any:
    if kind == "filename" and (value is None or not str(value).strip()):
        return "-"
    if value is None:
        return None
    if kind == "percent":
        return _percent_value(value)
    if kind == "integer":
        try:
            return int(value)
        except (TypeError, ValueError):
            return _safe_text(value)
    if kind in {"decimal", "money"}:
        try:
            return float(value)
        except (TypeError, ValueError):
            return _safe_text(value)
    if kind in {"date", "datetime"} and isinstance(value, (date, datetime)):
        return value
    return _safe_text(value)


def _make_sheet(workbook: Workbook, title: str, columns: Sequence[ExportColumn]):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 30
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(column.width, 40)
    headers = []
    for column in columns:
        cell = WriteOnlyCell(sheet, value=column.label)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        headers.append(cell)
    sheet.append(headers)
    return sheet


def _append_row(sheet, columns: Sequence[ExportColumn], item: Any) -> None:
    cells = []
    for column in columns:
        cell = WriteOnlyCell(sheet, value=_cell_value(column.getter(item), column.kind))
        cell.font = Font(name="微软雅黑", size=10, color="1F2937")
        cell.alignment = Alignment(
            horizontal="right" if column.kind in {"integer", "percent", "decimal", "money"} else "left",
            vertical="top",
            wrap_text=column.kind not in {"integer", "percent", "datetime"},
        )
        if column.kind == "datetime":
            cell.number_format = "yyyy-mm-dd hh:mm"
        elif column.kind == "date":
            cell.number_format = "yyyy-mm-dd"
        elif column.kind == "integer":
            cell.number_format = "#,##0"
        elif column.kind == "percent":
            cell.number_format = "0%"
        elif column.kind == "decimal":
            cell.number_format = "#,##0.0000"
        elif column.kind == "money":
            cell.number_format = "#,##0.00"
        elif column.kind == "identifier":
            cell.number_format = "@"
        cells.append(cell)
    sheet.append(cells)


